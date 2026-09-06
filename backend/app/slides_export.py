import io
import html
import re
from typing import List, Dict, Any, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

COLOR_BG = RGBColor(18, 18, 24)
COLOR_CARD_BG = RGBColor(28, 28, 38)
COLOR_ACCENT = RGBColor(139, 92, 246)
COLOR_CYAN = RGBColor(56, 189, 248)
COLOR_AMBER = RGBColor(245, 158, 11)
COLOR_TEXT_PRIMARY = RGBColor(241, 245, 249)
COLOR_TEXT_SECONDARY = RGBColor(203, 213, 225)
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)

SLIDE_TYPE_LABELS = {
    "title": "OVERVIEW",
    "concept": "CORE CONCEPT",
    "math": "MATHEMATICAL FORMALISM",
    "mechanism": "MECHANISM & DYNAMICS",
    "applications": "APPLICATIONS & FRONTIERS",
    "pitfalls": "PITFALLS & EDGE CASES",
    "summary": "SYNTHESIS & SELF-CHECK",
}

def clean_formula_for_export(formula: Optional[str]) -> Optional[str]:
    """Strip delimiters and clean formula for export to avoid double-delimited KaTeX errors."""
    if not formula:
        return None
    f = str(formula).strip()
    if f.lower() in ("none", "null", "n/a", "undefined", ""):
        return None
    if f.startswith("\\[") and f.endswith("\\]"):
        f = f[2:-2].strip()
    elif f.startswith("$$") and f.endswith("$$"):
        f = f[2:-2].strip()
    elif f.startswith("\\(") and f.endswith("\\)"):
        f = f[2:-2].strip()
    elif f.startswith("$") and f.endswith("$"):
        f = f[1:-1].strip()
    f = re.sub(r"^(\$\$|\$)+", "", f)
    f = re.sub(r"(\$\$|\$)+$", "", f).strip()
    return f if f else None

def build_pptx_deck(deck_title: str, slides: List[Dict[str, Any]]) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    total_slides = len(slides)

    for idx, slide_data in enumerate(slides):
        slide = prs.slides.add_slide(blank_layout)

        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

        title_text = slide_data.get("title", f"Slide {idx + 1}")
        slide_type = slide_data.get("slide_type", "concept")
        type_label = SLIDE_TYPE_LABELS.get(slide_type, slide_type.upper())
        subtitle = slide_data.get("subtitle", "")
        bullets = slide_data.get("bullets", [])
        callout = slide_data.get("callout")
        formula = clean_formula_for_export(slide_data.get("formula"))
        speaker_notes = slide_data.get("speaker_notes", "")
        quick_check = slide_data.get("quick_check")

        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.4))
        htf = header_box.text_frame
        htf.word_wrap = True
        htf.margin_left = htf.margin_top = htf.margin_right = htf.margin_bottom = 0

        p0 = htf.paragraphs[0]
        p0.text = f"● {type_label}"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_CYAN if slide_type in ["math", "mechanism"] else COLOR_ACCENT

        p1 = htf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_PRIMARY
        p1.space_before = Pt(4)

        if subtitle:
            p2 = htf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(13)
            p2.font.color.rgb = COLOR_TEXT_MUTED
            p2.space_before = Pt(2)

        has_right_panel = bool(formula or callout or quick_check)
        main_width = Inches(7.5) if has_right_panel else Inches(11.7)

        bullets_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.1), main_width, Inches(4.5))
        btf = bullets_box.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0

        for b_idx, bullet in enumerate(bullets):
            p = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.size = Pt(16)
            p.font.color.rgb = COLOR_TEXT_SECONDARY
            p.space_after = Pt(14)
            p.line_spacing = 1.25

        if has_right_panel:
            right_left = Inches(8.7)
            right_width = Inches(3.8)
            current_top = Inches(2.1)

            if formula:
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_left, current_top, right_width, Inches(1.8))
                card.fill.solid()
                card.fill.fore_color.rgb = COLOR_CARD_BG
                card.line.color.rgb = COLOR_ACCENT
                card.line.width = Pt(1.5)

                ctf = card.text_frame
                ctf.word_wrap = True
                ctf.margin_left = ctf.margin_right = Inches(0.2)
                ctf.margin_top = Inches(0.15)

                cp0 = ctf.paragraphs[0]
                cp0.text = "KEY GOVERNING FORMULA"
                cp0.font.size = Pt(10)
                cp0.font.bold = True
                cp0.font.color.rgb = COLOR_ACCENT

                cp1 = ctf.add_paragraph()
                cp1.text = formula
                cp1.font.size = Pt(14)
                cp1.font.bold = True
                cp1.font.color.rgb = COLOR_TEXT_PRIMARY
                cp1.space_before = Pt(8)

                current_top += Inches(2.0)

            if callout:
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_left, current_top, right_width, Inches(1.8))
                card.fill.solid()
                card.fill.fore_color.rgb = COLOR_CARD_BG
                card.line.color.rgb = COLOR_AMBER
                card.line.width = Pt(1.5)

                ctf = card.text_frame
                ctf.word_wrap = True
                ctf.margin_left = ctf.margin_right = Inches(0.2)
                ctf.margin_top = Inches(0.15)

                cp0 = ctf.paragraphs[0]
                cp0.text = "INSIGHT & PRINCIPLE"
                cp0.font.size = Pt(10)
                cp0.font.bold = True
                cp0.font.color.rgb = COLOR_AMBER

                cp1 = ctf.add_paragraph()
                cp1.text = callout
                cp1.font.size = Pt(12)
                cp1.font.color.rgb = COLOR_TEXT_SECONDARY
                cp1.space_before = Pt(6)

                current_top += Inches(2.0)

            if quick_check and not callout:
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_left, current_top, right_width, Inches(1.8))
                card.fill.solid()
                card.fill.fore_color.rgb = COLOR_CARD_BG
                card.line.color.rgb = COLOR_CYAN
                card.line.width = Pt(1.5)

                ctf = card.text_frame
                ctf.word_wrap = True
                ctf.margin_left = ctf.margin_right = Inches(0.2)
                ctf.margin_top = Inches(0.15)

                cp0 = ctf.paragraphs[0]
                cp0.text = "ACTIVE RECALL CHECK"
                cp0.font.size = Pt(10)
                cp0.font.bold = True
                cp0.font.color.rgb = COLOR_CYAN

                cp1 = ctf.add_paragraph()
                cp1.text = quick_check
                cp1.font.size = Pt(12)
                cp1.font.color.rgb = COLOR_TEXT_SECONDARY
                cp1.space_before = Pt(6)

        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
        ftf = footer_box.text_frame
        ftf.margin_left = ftf.margin_right = ftf.margin_top = ftf.margin_bottom = 0
        fp = ftf.paragraphs[0]
        fp.text = f"{deck_title}  |  Slide {idx + 1} of {total_slides}"
        fp.font.size = Pt(10)
        fp.font.color.rgb = COLOR_TEXT_MUTED

        if speaker_notes:
            notes_slide = slide.notes_slide
            notes_tf = notes_slide.notes_text_frame
            notes_tf.text = f"STUDY NOTES & NARRATION:\n\n{speaker_notes}"
            if quick_check:
                notes_tf.text += f"\n\nSELF-CHECK PROMPT:\n{quick_check}"

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def build_html_presentation(deck_title: str, slides: List[Dict[str, Any]]) -> str:
    total_slides = len(slides)

    slides_html_list = []
    for idx, s in enumerate(slides):
        title = html.escape(s.get("title", f"Slide {idx + 1}"))
        slide_type = s.get("slide_type", "concept")
        type_label = SLIDE_TYPE_LABELS.get(slide_type, slide_type.upper())
        subtitle = html.escape(s.get("subtitle", ""))
        bullets = s.get("bullets", [])
        callout = html.escape(s.get("callout", "")) if s.get("callout") else None
        raw_formula = clean_formula_for_export(s.get("formula"))
        formula = html.escape(raw_formula) if raw_formula else None
        speaker_notes = html.escape(s.get("speaker_notes", ""))
        quick_check = html.escape(s.get("quick_check", "")) if s.get("quick_check") else None

        bullets_li = "".join([f"<li>{html.escape(b)}</li>" for b in bullets])

        right_panel = ""
        if formula or callout or quick_check:
            right_panel = '<div class="slide-right-panel">'
            if formula:
                right_panel += f'<div class="card formula-card"><div class="card-header">KEY FORMULA</div><div class="card-body formula-text">$${formula}$$</div></div>'
            if callout:
                right_panel += f'<div class="card callout-card"><div class="card-header">KEY INSIGHT</div><div class="card-body">{callout}</div></div>'
            if quick_check:
                right_panel += f'<div class="card check-card"><div class="card-header">ACTIVE RECALL</div><div class="card-body">{quick_check}</div></div>'
            right_panel += '</div>'

        subtitle_html = f'<div class="slide-subtitle">{subtitle}</div>' if subtitle else ""
        notes_html = f'<div class="study-notes-drawer" id="notes-{idx}"><div class="notes-header">In-Depth Study Notes</div><div class="notes-body">{speaker_notes}</div></div>' if speaker_notes else ""

        active_cls = "active" if idx == 0 else ""
        has_right_cls = "has-right" if right_panel else ""

        slide_block = f"""
        <div class="slide-page {active_cls}" id="slide-{idx}" data-index="{idx}">
            <div class="slide-header">
                <span class="slide-type-badge badge-{slide_type}">{type_label}</span>
                <h1 class="slide-title">{title}</h1>
                {subtitle_html}
            </div>
            <div class="slide-body {has_right_cls}">
                <div class="slide-left-content">
                    <ul class="bullet-list">
                        {bullets_li}
                    </ul>
                </div>
                {right_panel}
            </div>
            <div class="slide-footer">
                <span class="footer-title">{html.escape(deck_title)}</span>
                <span class="footer-counter">Slide {idx + 1} of {total_slides}</span>
            </div>
            {notes_html}
        </div>"""
        slides_html_list.append(slide_block)

    all_slides_html = "\n".join(slides_html_list)

    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{html.escape(deck_title)} - Study Slides</title>
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/katex@0.16.8/dist/katex.min.css">
    <script defer src="https://cdn.jsdelivr.net/npm/katex@0.16.8/dist/katex.min.js"></script>
    <script defer src="https://cdn.jsdelivr.net/npm/katex@0.16.8/dist/contrib/auto-render.min.js"></script>
    <style>
        :root {{
            --bg-canvas: #0f0f14;
            --bg-slide: #16161e;
            --bg-card: #1f1f2e;
            --accent: #8b5cf6;
            --accent-light: #a78bfa;
            --cyan: #38bdf8;
            --amber: #f59e0b;
            --text-primary: #f1f5f9;
            --text-secondary: #cbd5e1;
            --text-muted: #94a3b8;
            --border-subtle: #2d2d3f;
        }}
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        body {{
            background-color: var(--bg-canvas);
            color: var(--text-primary);
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            height: 100vh;
            display: flex;
            flex-direction: column;
            overflow: hidden;
            user-select: none;
        }}
        .top-navbar {{
            display: flex;
            align-items: center;
            justify-content: space-between;
            padding: 10px 24px;
            background: rgba(22, 22, 30, 0.95);
            border-bottom: 1px solid var(--border-subtle);
            z-index: 10;
        }}
        .deck-title {{
            font-size: 15px;
            font-weight: 600;
            color: var(--text-primary);
            overflow: hidden;
            text-overflow: ellipsis;
            white-space: nowrap;
            max-width: 50vw;
        }}
        .top-controls {{
            display: flex;
            align-items: center;
            gap: 12px;
        }}
        button.nav-btn {{
            background: var(--bg-card);
            color: var(--text-primary);
            border: 1px solid var(--border-subtle);
            border-radius: 6px;
            padding: 6px 14px;
            font-size: 13px;
            cursor: pointer;
            transition: all 0.2s;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }}
        button.nav-btn:hover {{
            background: var(--border-subtle);
            border-color: var(--accent);
        }}
        .slide-stage {{
            flex: 1;
            display: flex;
            align-items: center;
            justify-content: center;
            padding: 24px;
            overflow: hidden;
            position: relative;
        }}
        .slide-page {{
            display: none;
            width: 100%;
            max-width: 1100px;
            aspect-ratio: 16 / 9;
            background: var(--bg-slide);
            border: 1px solid var(--border-subtle);
            border-radius: 12px;
            box-shadow: 0 20px 40px rgba(0, 0, 0, 0.6);
            padding: 36px 44px;
            position: relative;
            flex-direction: column;
            justify-content: space-between;
        }}
        .slide-page.active {{
            display: flex;
            animation: fadeIn 0.25s ease-out;
        }}
        @keyframes fadeIn {{
            from {{ opacity: 0; transform: translateY(6px); }}
            to {{ opacity: 1; transform: translateY(0); }}
        }}
        .slide-type-badge {{
            display: inline-block;
            font-size: 11px;
            font-weight: 700;
            letter-spacing: 0.8px;
            text-transform: uppercase;
            padding: 3px 8px;
            border-radius: 4px;
            margin-bottom: 8px;
            background: rgba(139, 92, 246, 0.15);
            color: var(--accent-light);
            border: 1px solid rgba(139, 92, 246, 0.3);
        }}
        .slide-title {{
            font-size: 28px;
            font-weight: 700;
            color: var(--text-primary);
            line-height: 1.25;
        }}
        .slide-subtitle {{
            font-size: 14px;
            color: var(--text-muted);
            margin-top: 4px;
        }}
        .slide-body {{
            flex: 1;
            display: flex;
            gap: 28px;
            margin: 20px 0;
            align-items: flex-start;
        }}
        .slide-left-content {{
            flex: 1;
        }}
        .bullet-list {{
            list-style: none;
            display: flex;
            flex-direction: column;
            gap: 14px;
        }}
        .bullet-list li {{
            font-size: 17px;
            line-height: 1.5;
            color: var(--text-secondary);
            position: relative;
            padding-left: 24px;
        }}
        .bullet-list li::before {{
            content: "•";
            position: absolute;
            left: 4px;
            color: var(--accent);
            font-size: 22px;
            line-height: 1;
        }}
        .slide-right-panel {{
            width: 380px;
            display: flex;
            flex-direction: column;
            gap: 14px;
        }}
        .card {{
            background: var(--bg-card);
            border: 1px solid var(--border-subtle);
            border-radius: 8px;
            padding: 14px 16px;
        }}
        .formula-card {{ border-color: rgba(56, 189, 248, 0.4); }}
        .callout-card {{ border-color: rgba(245, 158, 11, 0.4); }}
        .check-card {{ border-color: rgba(139, 92, 246, 0.4); }}
        .card-header {{
            font-size: 11px;
            font-weight: 700;
            letter-spacing: 0.6px;
            margin-bottom: 6px;
            color: var(--text-muted);
        }}
        .formula-text {{
            font-size: 15px;
            color: var(--cyan);
            overflow-x: auto;
            text-align: center;
        }}
        .slide-footer {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            font-size: 12px;
            color: var(--text-muted);
            border-top: 1px solid var(--border-subtle);
            padding-top: 12px;
        }}
        .study-notes-drawer {{
            margin-top: 8px;
            background: rgba(0, 0, 0, 0.35);
            border: 1px dashed var(--border-subtle);
            border-radius: 6px;
            padding: 8px 12px;
            font-size: 12px;
            color: var(--text-muted);
            line-height: 1.4;
        }}
        .notes-header {{
            font-weight: 600;
            color: var(--accent-light);
            margin-bottom: 2px;
            font-size: 11px;
        }}
        .progress-bar {{
            height: 3px;
            background: var(--border-subtle);
            width: 100%;
        }}
        .progress-fill {{
            height: 100%;
            background: var(--accent);
            width: 14.28%;
            transition: width 0.2s;
        }}
        @media print {{
            body {{ background: #fff; color: #000; overflow: visible; height: auto; }}
            .top-navbar, .progress-bar, .nav-btn {{ display: none; }}
            .slide-stage {{ padding: 0; display: block; }}
            .slide-page {{
                display: flex !important;
                page-break-after: always;
                box-shadow: none;
                border: 1px solid #ddd;
                background: #fff;
                color: #000;
                margin-bottom: 20px;
            }}
            .slide-title {{ color: #111; }}
            .bullet-list li {{ color: #333; }}
        }}
    </style>
</head>
<body>
    <div class="progress-bar">
        <div class="progress-fill" id="progressFill"></div>
    </div>

    <header class="top-navbar">
        <div class="deck-title">{html.escape(deck_title)}</div>
        <div class="top-controls">
            <button class="nav-btn" onclick="prevSlide()" title="Previous slide (Left Arrow)">◀ Prev</button>
            <span id="slideIndicator" style="font-size: 13px; color: var(--text-muted);">1 / {total_slides}</span>
            <button class="nav-btn" onclick="nextSlide()" title="Next slide (Right Arrow / Space)">Next ▶</button>
            <button class="nav-btn" onclick="window.print()" title="Print or Save as PDF">Print / PDF</button>
        </div>
    </header>

    <main class="slide-stage">
        {all_slides_html}
    </main>

    <script>
        let currentIdx = 0;
        const total = {total_slides};

        function showSlide(idx) {{
            if (idx < 0 || idx >= total) return;
            document.querySelectorAll('.slide-page').forEach(el => el.classList.remove('active'));
            const target = document.getElementById('slide-' + idx);
            if (target) target.classList.add('active');
            currentIdx = idx;

            document.getElementById('slideIndicator').innerText = (currentIdx + 1) + ' / ' + total;
            document.getElementById('progressFill').style.width = (((currentIdx + 1) / total) * 100) + '%';
        }}

        function nextSlide() {{
            if (currentIdx < total - 1) showSlide(currentIdx + 1);
        }}

        function prevSlide() {{
            if (currentIdx > 0) showSlide(currentIdx - 1);
        }}

        window.addEventListener('keydown', (e) => {{
            if (e.key === 'ArrowRight' || e.key === ' ' || e.key === 'PageDown') {{
                e.preventDefault();
                nextSlide();
            }} else if (e.key === 'ArrowLeft' || e.key === 'PageUp') {{
                e.preventDefault();
                prevSlide();
            }}
        }});

        window.addEventListener('DOMContentLoaded', () => {{
            if (typeof renderMathInElement === 'function') {{
                renderMathInElement(document.body, {{
                    delimiters: [
                        {{left: '$$', right: '$$', display: true}},
                        {{left: '$', right: '$', display: false}}
                    ],
                    throwOnError: false
                }});
            }}
        }});
    </script>
</body>
</html>
"""
    return html_template
